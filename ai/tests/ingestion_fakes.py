"""Test documents built in code (no binary fixtures) and a scripted OCR engine."""

import io
import textwrap
from collections.abc import Sequence
from pathlib import Path
from typing import Any

import docx
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFont

from legallens_ai.ingestion.ocr import OcrLine

DPI = 200
LETTER = (int(8.5 * DPI), int(11 * DPI))


def scanned_image(paragraphs: Sequence[str] = (), *, font_px: int = 34) -> Image.Image:
    """A letter-size page image of the text, like a clean 200 DPI scan (blank if no text)."""
    image = Image.new("L", LETTER, 255)
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=font_px)
    y = DPI
    for paragraph in paragraphs:
        for line in textwrap.wrap(paragraph, 70):
            draw.text((DPI, y), line, font=font, fill=0)
            y += int(font_px * 1.35)
        y += int(font_px * 0.9)
    return image


Page = list[str] | Image.Image


def make_pdf(path: Path, pages: Sequence[Page]) -> Path:
    """Each page is either text (a real text layer) or an image (no text layer, like a scan)."""
    pdf = FPDF(unit="pt", format="letter")
    pdf.set_font("helvetica", size=11)
    for page in pages:
        pdf.add_page()
        if isinstance(page, Image.Image):
            pdf.image(page, x=0, y=0, w=pdf.w, h=pdf.h)
        else:
            for paragraph in page:
                pdf.multi_cell(0, 14, paragraph, new_x="LMARGIN", new_y="NEXT")
                pdf.ln(8)
    path.write_bytes(bytes(pdf.output()))
    return path


def make_docx(
    path: Path, pages: Sequence[Sequence[str]], *, table: list[list[str]] | None = None
) -> Path:
    document = docx.Document()
    for number, paragraphs in enumerate(pages):
        if number:
            document.add_page_break()  # type: ignore[no-untyped-call]
        for paragraph in paragraphs:
            if paragraph.startswith("#"):
                document.add_heading(paragraph.lstrip("# "), level=1)
            else:
                document.add_paragraph(paragraph)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    out = io.BytesIO()
    document.save(out)
    path.write_bytes(out.getvalue())
    return path


def line(text: str, top: float, *, left: float = 200, score: float = 0.98) -> OcrLine:
    """An OCR line 30 px tall whose width follows its length."""
    return OcrLine(text, (left, top, left + 18 * len(text), top + 30), score)


class ScriptedOcr:
    """Returns the next scripted page of lines on each call; raises ``error`` if set."""

    name = "scripted"

    def __init__(self, pages: Sequence[list[OcrLine]] = (), error: Exception | None = None) -> None:
        self.pages = list(pages)
        self.error = error
        self.images: list[tuple[int, ...]] = []

    async def read(self, image: Any) -> list[OcrLine]:
        self.images.append(tuple(image.shape))
        if self.error:
            raise self.error
        return self.pages.pop(0) if self.pages else []
