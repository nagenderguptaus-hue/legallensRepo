"""Builds test documents in code, so no binary fixtures are checked in."""

import io
from datetime import UTC, datetime

import docx
from fpdf import FPDF
from pypdf import PdfWriter

FIXED_DATE = datetime(2026, 1, 1, tzinfo=UTC)

# A small synthetic NDA, one list of paragraphs per page.
NDA_PAGES: list[list[str]] = [
    [
        "MUTUAL NON-DISCLOSURE AGREEMENT",
        "This Mutual Non-Disclosure Agreement is entered into on March 1, 2025 by and between "
        "Acme Corp and Beta LLC.",
        "1. Definitions",
        "Confidential Information means any non-public information disclosed by either party.",
    ],
    [
        "2. Term",
        "This Agreement shall remain in effect for a period of three (3) years from the "
        "Effective Date.",
        "3. Obligations",
        "The Receiving Party shall protect Confidential Information with reasonable care.",
    ],
    [
        "4. Termination",
        "Either party may terminate this Agreement with thirty (30) days written notice.",
        "The Receiving Party shall be liable for any breach of this Agreement.",
        "5. Governing Law",
        "This Agreement is governed by the laws of the State of California.",
    ],
]


def make_pdf(pages: list[list[str]] | None = None) -> bytes:
    pdf = FPDF()
    # Fixed timestamp: same input, same bytes, same SHA-256 (duplicate detection is by hash).
    pdf.set_creation_date(FIXED_DATE)
    pdf.set_font("helvetica", size=11)
    for paragraphs in pages or NDA_PAGES:
        pdf.add_page()
        for paragraph in paragraphs:
            pdf.multi_cell(0, 6, paragraph, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(4)
    return bytes(pdf.output())


def make_blank_pdf(pages: int = 2) -> bytes:
    """A PDF with no text layer, like a scan."""
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=612, height=792)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_encrypted_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt(user_password="secret", owner_password="owner")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx(pages: list[list[str]] | None = None, *, page_breaks: bool = True) -> bytes:
    document = docx.Document()
    for number, paragraphs in enumerate(pages or NDA_PAGES):
        if number and page_breaks:
            document.add_page_break()  # type: ignore[no-untyped-call]
        for paragraph in paragraphs:
            if paragraph.isupper() or paragraph[:1].isdigit():
                document.add_heading(paragraph, level=1)
            else:
                document.add_paragraph(paragraph)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
